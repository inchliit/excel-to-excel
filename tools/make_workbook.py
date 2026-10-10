"""Build downloads/Excel-to-Excel-practice.xlsx from the deck's own example data (tools/examples.json).

Each trick gets one sheet. The data sits at the SAME cell addresses as on the slide, so every formula shown in the
deck works verbatim. Instructions go to the right of the data; yellow cells are the student's; a Check cell gives
instant feedback (classic functions only, so it works in every Excel version); answers sit in a collapsed column group.
Also writes tools/workbook-verify.json for tests/check-workbook.mjs (HyperFormula recalculation + solution checks).

Regenerate examples.json from the deck first if slides changed (see tools/README-workbook.md), then:
    python3 tools/make_workbook.py
"""
import datetime
import json
import math
import pathlib
import re
import textwrap

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
DATA = json.loads((ROOT / 'tools/examples.json').read_text(encoding='utf-8'))
OUT = ROOT / 'downloads/Excel-to-Excel-practice.xlsx'
VERIFY = ROOT / 'tools/workbook-verify.json'
DECK_URL = 'https://inchliit.github.io/excel-to-excel/'

FONT = 'Arial'
YELLOW = PatternFill('solid', fgColor='FFFF00')
HEAD = PatternFill('solid', fgColor='E2EFDA')
SOFT = PatternFill('solid', fgColor='FFF9D6')
GREEN = '107C41'
THIN = Side(style='thin', color='B7B7B7')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
START = 'Start here'
VERSION_TEXT = {
    'All': 'Works in Excel 2013 and newer',
    '2016+': 'Works in Excel 2016, 2019, 2021 and Microsoft 365',
    '2019+': 'Works in Excel 2019, 2021 and Microsoft 365',
    '365/2021': 'Needs Microsoft 365 or Excel 2021 (older Excel shows #NAME?)',
}
SHEETS = {1: '01 Flash Fill', 2: '02 Go To Special', 3: '03 Paste Special', 4: '04 Tables', 5: '05 Data Validation',
          6: '06 Mixed Refs', 7: '07 SUMIFS', 8: '08 INDEX MATCH', 9: '09 XLOOKUP', 10: '10 IFS', 11: '11 FILTER SORT',
          12: '12 Cond Formatting', 13: '13 Sparklines', 14: '14 Pivot Slicers', 15: '15 Pareto', 16: '16 Goal Seek',
          17: '17 Data Table', 18: '18 Solver', 19: '19 Forecast Sheet', 20: '20 Monte Carlo'}
FMT = {'0': '0', '0.0': '0.0', '0.00': '0.00', '0.000': '0.000', '0%': '0%', '0.0%': '0.0%', '#,##0': '#,##0',
       '#,##0.00': '#,##0.00', 'text': '@'}


# ---------- references ----------
def rc(ref):
    m = re.match(r'^\$?([A-Z]+)\$?(\d+)$', ref)
    return int(m.group(2)), column_index_from_string(m.group(1))


def bounds(rng):
    a, _, b = rng.partition(':')
    (r0, c0), (r1, c1) = rc(a), rc(b or a)
    return min(r0, r1), min(c0, c1), max(r0, r1), max(c0, c1)


def a1(r, c):
    return f'{get_column_letter(c)}{r}'


def cells(rng):
    r0, c0, r1, c1 = bounds(rng)
    return [a1(r, c) for r in range(r0, r1 + 1) for c in range(c0, c1 + 1)]


def shift(rng, dr, dc):
    r0, c0, r1, c1 = bounds(rng)
    return f'{a1(r0 + dr, c0 + dc)}:{a1(r1 + dr, c1 + dc)}'


# ---------- per-trick practice design (data comes from the deck; this adds the task, steps and checks) ----------
def final_block(t, rng, grid=0):
    vals = t['final'][grid]['vals']
    r0, c0, r1, c1 = bounds(rng)
    return [[vals.get(a1(r, c)) for c in range(c0, c1 + 1)] for r in range(r0, r1 + 1)]


HIST = [round(1000 + 12 * t + 180 * math.sin(2 * math.pi * t / 12) + 20 * math.sin(2.4 * t)) for t in range(1, 25)]

DESIGN = {
    'flash-fill': dict(src='initial', inputs=['B2:D9'],
        steps=['Type L2 in B2 (the line part of A2).', 'Click B3 and press Ctrl+E: Excel fills the rest of the column.',
               'Do the same for Operator (type OP0423 in C2) and Shift (type Night in D2), then Ctrl+E.',
               'Also on the ribbon: Data › Flash Fill.'],
        check=dict(kind='exact', range='B2:D9', expected='final')),
    'go-to-special': dict(src='initial', inputs=['A3', 'A4', 'A6', 'A7', 'A9'],
        steps=['Select A2:A9.', 'Press F5 (or Ctrl+G) › Special… › Blanks › OK. Only the empty cells stay selected.',
               'Type = then press ↑ (the formula bar shows =A2).', 'Press Ctrl+Enter: every blank copies the cell above.',
               'Optional: copy A2:A9 and paste as Values (Ctrl+Alt+V › Values) to freeze it.'],
        check=dict(kind='exact', range='A2:A9', expected='final')),
    'paste-special-math': dict(src='initial', inputs=['B2:E6', 'G2'],
        steps=['Type 60 in G2 and copy it (Ctrl+C).', 'Select B2:E6 (the seconds).',
               'Press Ctrl+Alt+V (Mac: Cmd+Ctrl+V) › Paste: Values › Operation: Divide › OK.',
               'The times are now minutes, in place. You can delete G2 afterwards.'],
        check=dict(kind='num', range='B2:E6', expected='final', digits=4)),
    'tables': dict(src='initial', inputs=['E2:E7'],
        steps=['Click any cell in A1:E7 and press Ctrl+T › OK (My table has headers).',
               'In E2 type =1-[@Defects]/[@Units] and press Enter: the whole column fills.',
               'Add a new day in row 8 (e.g. 8-Oct, L1, 455, 7): the table and the Yield formula grow by themselves.',
               'Bonus: Ctrl+Shift+T adds a Total Row.'],
        check=dict(kind='num', range='E2:E7', expected='final', digits=4)),
    'data-validation': dict(src='initial', inputs=['D2:D11'],
        steps=['Select D2:D11.', 'Data › Data Tools › Data Validation (keys: Alt, A, V, V).',
               'Allow: List · Source: =$H$2:$H$6 · tick In-cell dropdown › OK.',
               'Click D2: pick from the list. Try typing crak: Excel refuses it.',
               'Data › Data Validation (arrow) › Circle Invalid Data shows the old bad entries.'],
        see='The typo cells (scrtch, Scrach, "Dent " with a trailing space, Crak) get red circles.'),
    'mixed-refs': dict(src='tryit', inputs=['C2:G6'], fill=True),
    'sumifs': dict(src='tryit', inputs=['G2'], labels={'F1': 'Question', 'F2': 'L2, Night', 'G1': 'Answer'}),
    'index-match': dict(src='tryit', inputs=['G4']),
    'xlookup': dict(src='tryit', inputs=['G2']),
    'ifs': dict(src='tryit', inputs=['D2:D11'], fill=True),
    'filter-sort': dict(src='tryit', inputs=['H2'], soft=['H2:M5'],
        labels={'H1': 'Machine', 'I1': 'Line', 'J1': 'Avail', 'K1': 'Perf', 'L1': 'Quality', 'M1': 'OEE'}),
    'cf-formula': dict(src='tryit', inputs=['H2:H21'], fill=True,
        steps_extra=['Then make it visual: select A2:B21 › Home › Conditional Formatting › New Rule ›',
                     '"Use a formula to determine which cells to format" › =OR($B2>$F$2,$B2<$F$3) › Format… › red fill › OK.',
                     'Change F2 (UCL) to 25.05 and watch one more row light up.']),
    'sparklines': dict(src='initial', inputs=['I2:I6'],
        steps=['Select B2:H6.', 'Insert › Sparklines › Line.', 'Location Range: $I$2:$I$6 › OK.',
               'Sparkline tab › Show: tick High Point and Low Point. Try Type › Column too.'],
        see='One mini chart per line; L4 slopes down all week (468 → 386).'),
    'pivot-slicers': dict(src='initial', inputs=[],
        steps=['Click any cell in A1:D17 › Insert › PivotTable › New Worksheet › OK.',
               'Drag Cause to Rows and Minutes to Values.',
               'Right-click a number › Sort › Largest to Smallest.',
               'PivotTable Analyze › Insert Slicer › tick Line › OK. Click L2, then L3.',
               'New data later? Alt+F5 refreshes this pivot, Ctrl+Alt+F5 refreshes everything.'],
        see='pivot'),
    'pareto': dict(src='initial', inputs=[],
        steps=['Select A1:B8.', 'Insert › Insert Statistic Chart › Pareto (under Histogram).',
               'Excel sorts the causes and adds the cumulative % line for you.'],
        see='pareto'),
    'goal-seek': dict(src='initial', inputs=['B2'],
        steps=['Data › What-If Analysis › Goal Seek.', 'Set cell: B3 · To value: 480 · By changing cell: B2 › OK.'],
        check=dict(kind='goal')),
    'data-table': dict(src='initial', inputs=['E8:E15'], setup='data-table',
        steps=['Select D7:E15 (Q values in D, the TC link =B5 in E7).',
               'Data › What-If Analysis › Data Table.', 'Row input cell: leave blank · Column input cell: B4 › OK.',
               'Excel fills E8:E15 with {=TABLE(,B4)}. The cheapest Q sits between 1,000 and 1,200 (EOQ ≈ 1,095).'],
        check=dict(kind='num', range='E8:E15', expected='final', digits=0)),
    'solver': dict(src='initial', inputs=['B2:C2'],
        steps=['First time only: File › Options › Add-ins › Manage: Excel Add-ins › Go › tick Solver Add-in › OK.',
               '(Mac: Tools › Excel Add-ins › Solver Add-In.)',
               'Data › Solver. Set Objective: $D$3 · To: Max · By Changing Variable Cells: $B$2:$C$2.',
               'Add the constraint $D$5:$D$6 <= $F$5:$F$6 · tick Make Unconstrained Variables Non-Negative.',
               'Solving Method: Simplex LP › Solve › Keep Solver Solution › OK.'],
        check=dict(kind='solver')),
    'forecast-sheet': dict(src='forecast', inputs=[],
        steps=['Click any cell in A1:B25 (one cell inside the data is enough).',
               'Data › Forecast › Forecast Sheet (Windows Excel 2016+).',
               'Check Options: Confidence Interval 95%, Seasonality: Detect Automatically › Create.',
               'Mac / web: no Forecast Sheet button. Type =FORECAST.ETS(…) yourself.'],
        see='A new sheet with a forecast line and a 95% confidence band for the next months.'),
    'monte-carlo': dict(src='tryit', inputs=['D3'],
        steps_extra=['Then fill D2 and D4 with their own means and SDs, and add =SUM(D2:D4) for the total.',
                     'Press F9 to re-roll. Fill 1,000 rows and COUNTIF the totals ≤ 42 to estimate on-time %.']),
}


def build():
    wb = Workbook()
    start = wb.active
    start.title = START
    verify = {'sheets': {}, 'checks': []}
    tricks = sorted(DATA.values(), key=lambda t: t['n'])
    for t in tricks:
        ws = wb.create_sheet(SHEETS[t['n']])
        info = make_trick_sheet(ws, t)
        verify['checks'].append(info)
    make_start(start, tricks)
    for ws in wb.worksheets:
        verify['sheets'][ws.title] = dump(ws)
    wb.calculation.fullCalcOnLoad = True
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    VERIFY.write_text(json.dumps(verify, ensure_ascii=False), encoding='utf-8')
    print('wrote', OUT.relative_to(ROOT), OUT.stat().st_size, 'bytes;', len(tricks), 'trick sheets')


def put(ws, ref, value, fmt=None, bold=False, fill=None, color=None, size=10, italic=False, align=None):
    c = ws[ref]
    c.value = value
    c.font = Font(name=FONT, bold=bold, size=size, color=color, italic=italic)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if align:
        c.alignment = align
    return c


def write_grid(ws, grid, header=True):
    for ref, v in grid['vals'].items():
        if v is None or v == '':
            continue
        f = grid['forms'].get(ref)
        fmt = None
        r, c = rc(ref)
        for rng, code in (grid.get('formats') or {}).items():
            r0, c0, r1, c1 = bounds(rng)
            if r0 <= r <= r1 and c0 <= c <= c1:
                fmt = FMT.get(code)
        put(ws, ref, f if f and not f.startswith('{') else v, fmt=fmt, bold=header and r == 1,
            fill=HEAD if header and r == 1 else None)
    for i, w in enumerate(grid.get('widths') or []):
        ws.column_dimensions[get_column_letter(i + 1)].width = max(9, round(w / 7, 1))


def write_rows(ws, rows, formats=None):
    for r, row in enumerate(rows, start=1):
        for c, v in enumerate(row, start=1):
            if v is None or v == '':
                continue
            put(ws, a1(r, c), v, bold=r == 1, fill=HEAD if r == 1 else None, fmt=(formats or {}).get(c) if r > 1 else None)
    for c in range(1, max(len(r) for r in rows) + 1):
        longest = max(len(str(row[c - 1])) if c - 1 < len(row) and row[c - 1] is not None else 0 for row in rows)
        ws.column_dimensions[get_column_letter(c)].width = max(10, min(22, longest + 3))


def used_cols(ws):
    return max((c.column for row in ws.iter_rows() for c in row if c.value not in (None, '')), default=1)


def make_trick_sheet(ws, t):
    d = DESIGN[t['id']]
    ti = t.get('tryit')
    # 1) data at the slide's own addresses
    if d['src'] == 'tryit':
        write_rows(ws, ti['sheet'])
        init = t['initial'][0]
        for rng, code in (init.get('formats') or {}).items():  # reuse the slide's number formats where they fit
            for ref in cells(rng):
                r, c = rc(ref)
                if r <= len(ti['sheet']) and c <= len(ti['sheet'][r - 1]) and ws[ref].value is not None:
                    ws[ref].number_format = FMT.get(code, ws[ref].number_format)
    elif d['src'] == 'forecast':
        put(ws, 'A1', 'Month', bold=True, fill=HEAD)
        put(ws, 'B1', 'Demand', bold=True, fill=HEAD)
        for i, v in enumerate(HIST):
            y, m = 2024 + (9 + i) // 12, (9 + i) % 12 + 1
            put(ws, f'A{i + 2}', datetime.date(y, m, 1), fmt='mmm-yy')
            put(ws, f'B{i + 2}', v, fmt='#,##0')
        ws.column_dimensions['A'].width = 11
        ws.column_dimensions['B'].width = 11
    else:
        write_grid(ws, t['initial'][0])
    for ref, text in (d.get('labels') or {}).items():
        put(ws, ref, text, bold=rc(ref)[0] == 1, fill=HEAD if rc(ref)[0] == 1 else None)
    if d.get('setup') == 'data-table':
        put(ws, 'D6', 'Q', bold=True, fill=HEAD)
        put(ws, 'E6', 'TC', bold=True, fill=HEAD)
        put(ws, 'E7', '=B5', fmt='#,##0')
        for i, q in enumerate(range(400, 1801, 200)):
            put(ws, f'D{8 + i}', q, fmt='#,##0')
        ws.column_dimensions['D'].width = 10
        ws.column_dimensions['E'].width = 12
    for rng in d.get('soft') or []:
        for ref in cells(rng):
            ws[ref].fill = SOFT
    for rng in d['inputs']:
        for ref in cells(rng):
            ws[ref].fill = YELLOW
            ws[ref].border = BOX

    # 2) instructions panel to the right of the data
    pc = used_cols(ws) + 2
    P = get_column_letter(pc)
    ans_c = pc + 8                      # answer group starts here (collapsed)
    A = get_column_letter(ans_c)
    row = 1
    put(ws, f'{P}{row}', f"{t['n']:02d} · {t['title']}", bold=True, size=14, color=GREEN); row += 1
    put(ws, f'{P}{row}', VERSION_TEXT[t['version']], italic=True, color='595959'); row += 1
    put(ws, f'{P}{row}', t['tagline'], bold=True); row += 1
    if t.get('alt'):
        put(ws, f'{P}{row}', 'Older Excel: ' + re.sub(r'^[^:]{3,30}:\s+', '', t['alt']), color='9C5700'); row += 1
    row += 1
    put(ws, f'{P}{row}', 'YOUR TURN  (yellow cells are yours)', bold=True, color='9C5700'); row += 1
    lines = []
    if d['src'] == 'tryit':
        lines.append(ti['task'])
    for i, s in enumerate(d.get('steps') or [], start=1):
        lines.append(f'{i}. {s}')
    lines += d.get('steps_extra') or []
    for text in lines:
        for k, part in enumerate(textwrap.wrap(text, 92)):
            put(ws, f'{P}{row}', ('    ' if k else '') + part); row += 1
    row += 1

    # 3) check cell + answers (collapsed group)
    check_ref, solution = None, None
    put(ws, f'{A}1', 'ANSWER', bold=True, color=GREEN)
    if ti:
        answer_note(ws, f'{A}2', ti['answer'])
        check_ref = f'{get_column_letter(pc + 1)}{row}'
        put(ws, f'{P}{row}', 'Check:', bold=True)
        exp_at = ans_c + 1
        if ti['mode'] == 'value':
            put(ws, f'{A}3', 'Expected')
            put(ws, f'{get_column_letter(exp_at)}3', ti['expected'])
            exp = f'${get_column_letter(exp_at)}$3'
            tgt = ti['target']
            ws[check_ref] = (f'=IF(ISBLANK({tgt}),"",IF(ISERROR({tgt}),"✗ Your formula returns an error",'
                             f'IF(ABS({tgt}-{exp})<0.000001,"✓ Correct!","✗ Not yet: check the ranges and criteria")))')
            solution = {'kind': 'formula', 'target': tgt, 'formula': ti['answer']}
        elif ti['mode'] in ('fill', 'cf', 'spill'):
            rng = ti['fill'] if ti['mode'] != 'spill' else spill_range(ti)
            r0, c0, r1, c1 = bounds(rng)
            block = ti['expected']
            put(ws, f'{A}3', 'Expected')
            for i, line in enumerate(block):
                for j, v in enumerate(line):
                    put(ws, a1(4 + i, exp_at + j), v)
            erng = f'${get_column_letter(exp_at)}$4:${get_column_letter(exp_at + c1 - c0)}${4 + r1 - r0}'
            n = (r1 - r0 + 1) * (c1 - c0 + 1)
            texty = any(isinstance(v, str) for line in block for v in line)
            boolish = all(isinstance(v, bool) for line in block for v in line)
            if texty:
                same = f'SUMPRODUCT(--EXACT({rng},{erng}))'
            elif boolish:
                same = f'SUMPRODUCT(--({rng}={erng}))'
            else:
                same = f'SUMPRODUCT(--(ROUND({rng},6)=ROUND({erng},6)))'
            ws[check_ref] = (f'=IF(COUNTBLANK({rng})={n},"",IF(SUMPRODUCT(--ISERROR({rng}))>0,"✗ Some cells show an error",'
                             f'IF({same}={n},"✓ Correct!","✗ Not yet: compare with the task")))')
            solution = {'kind': 'spill-values' if ti['mode'] == 'spill' else 'formula-fill', 'target': ti['target'],
                        'range': rng, 'formula': ti['answer'], 'values': block}
        elif ti['mode'] == 'random':
            tgt = ti['target']
            lo, hi = ti['random']['mean'] - 4 * ti['random']['sd'], ti['random']['mean'] + 4 * ti['random']['sd']
            ws[check_ref] = (f'=IF(ISBLANK({tgt}),"",IF(ISERROR({tgt}),"✗ Your formula returns an error",'
                             f'IF(AND(ISNUMBER({tgt}),{tgt}>{lo:g},{tgt}<{hi:g}),"✓ Looks right! Press F9 to draw a new time",'
                             f'"✗ Expect a time near {ti["random"]["mean"]:g} min")))')
            solution = {'kind': 'formula', 'target': tgt, 'formula': ti['answer']}
        put(ws, f'{A}{len(ti["expected"]) + 6 if isinstance(ti.get("expected"), list) else 6}', 'Hint: ' + ti['hint'], italic=True)
    elif d.get('check'):
        ck = d['check']
        check_ref = f'{get_column_letter(pc + 1)}{row}'
        put(ws, f'{P}{row}', 'Check:', bold=True)
        if ck['kind'] in ('exact', 'num'):
            block = final_block(t, ck['range'])
            r0, c0, r1, c1 = bounds(ck['range'])
            exp_at = ans_c + 1
            put(ws, f'{A}3', 'Expected (after the trick)')
            for i, line in enumerate(block):
                for j, v in enumerate(line):
                    put(ws, a1(4 + i, exp_at + j), v)
            erng = f'${get_column_letter(exp_at)}$4:${get_column_letter(exp_at + c1 - c0)}${4 + r1 - r0}'
            n = (r1 - r0 + 1) * (c1 - c0 + 1)
            rng = ck['range']
            same = (f'SUMPRODUCT(--EXACT({rng},{erng}))' if ck['kind'] == 'exact'
                    else f'SUMPRODUCT(--(ROUND({rng},{ck["digits"]})=ROUND({erng},{ck["digits"]})))')
            core = (f'IF(SUMPRODUCT(--ISERROR({rng}))>0,"✗ Some cells show an error",'
                    f'IF({same}={n},"✓ Done right!","✗ Not yet: follow the steps"))')
            # where the data is there from the start (Go To Special, Paste Special) there is no "untouched = blank" state
            ws[check_ref] = ('=' + core if t['id'] in ('go-to-special', 'paste-special-math')
                             else f'=IF(COUNTBLANK({rng})={n},"",{core})')
            solution = {'kind': 'values', 'range': rng, 'values': block}
            answer_note(ws, f'{A}2', ' '.join(d['steps'][:3]))
        elif ck['kind'] == 'goal':
            ws[check_ref] = ('=IF(ABS(B3-480)<0.01,"✓ Goal Seek done: cycle time "&TEXT(B2,"0.00")&" s = takt time",'
                             '"Run Goal Seek: Set cell B3, To value 480, By changing cell B2")')
            put(ws, f'{A}2', 'B2 becomes 56.25 s (27,000 s ÷ 480 units).')
            solution = {'kind': 'values', 'range': 'B2', 'values': [[56.25]]}
        elif ck['kind'] == 'solver':
            ws[check_ref] = ('=IF(AND(ROUND(B2,4)=20,ROUND(C2,4)=60),"✓ Optimal: 20 brackets + 60 housings = 2,600 profit",'
                             '"Run Solver: maximize D3 by changing B2:C2 (steps on the left)")')
            put(ws, f'{A}2', 'Optimum: B2 = 20 brackets, C2 = 60 housings, profit D3 = 2,600; both constraints 100% used.')
            solution = {'kind': 'values', 'range': 'B2:C2', 'values': [[20, 60]]}
    else:
        see = d.get('see', '')
        if see == 'pivot':
            see = pivot_text(t)
        elif see == 'pareto':
            see = pareto_text(t)
        put(ws, f'{P}{row}', 'What you should see:', bold=True)
        for k, part in enumerate(textwrap.wrap(see, 92)):
            put(ws, f'{P}{row + 1 + k}', part)
        put(ws, f'{A}2', see)
    if check_ref:
        ws[check_ref].font = Font(name=FONT, bold=True, size=11)
        ws.conditional_formatting.add(check_ref, FormulaRule(formula=[f'LEFT({check_ref},1)="✓"'], font=Font(name=FONT, bold=True, color='008000')))
        ws.conditional_formatting.add(check_ref, FormulaRule(formula=[f'LEFT({check_ref},1)="✗"'], font=Font(name=FONT, bold=True, color='C00000')))
        put(ws, f'{P}{row + 1}', f'Stuck? The answer is in hidden columns {A}–{get_column_letter(ans_c + 6)}: click the [+] in the bar above the column letters.',
            italic=True, color='595959')
        last = row + 1
    else:
        last = row + 3
    put(ws, f'{P}{last + 2}', f'=HYPERLINK("#\'{START}\'!A1","← Back to Start here")', color='0563C1')
    put(ws, f'{P}{last + 3}', f'=HYPERLINK("{DECK_URL}","Open the webinar deck")', color='0563C1')
    # collapse the answer columns into an outline group
    ws.column_dimensions.group(A, get_column_letter(ans_c + 6), hidden=True, outline_level=1)
    ws.sheet_properties.outlinePr.summaryRight = True
    ws.sheet_view.zoomScale = 110
    return {'sheet': ws.title, 'id': t['id'], 'check': check_ref, 'solution': solution}


def answer_note(ws, ref, text):
    c = put(ws, ref, text, bold=True)
    if isinstance(text, str) and text.startswith('='):
        c.data_type = 's'      # show the formula as text, don't run it
    return c


def spill_range(ti):
    r, c = rc(ti['target'])
    h, w = len(ti['expected']), len(ti['expected'][0])
    return f'{a1(r, c)}:{a1(r + h - 1, c + w - 1)}'


def pivot_text(t):
    vals = t['initial'][0]['vals']
    totals = {}
    for r in range(2, 18):
        cause, mins = vals.get(f'C{r}'), vals.get(f'D{r}')
        if cause and isinstance(mins, (int, float)):
            totals[cause] = totals.get(cause, 0) + mins
    parts = ', '.join(f'{k} {v:g}' for k, v in sorted(totals.items(), key=lambda kv: -kv[1]))
    return f'Sum of Minutes by Cause, largest first: {parts}; Grand Total {sum(totals.values()):g}. Click L2 / L3 on the slicer to filter.'


def pareto_text(t):
    vals = t['initial'][0]['vals']
    rows = [(vals.get(f'A{r}'), vals.get(f'B{r}')) for r in range(2, 9) if isinstance(vals.get(f'B{r}'), (int, float))]
    rows.sort(key=lambda x: -x[1])
    total = sum(v for _, v in rows)
    run, cum = 0, []
    for name, v in rows[:3]:
        run += v
        cum.append(f'{name} {run / total:.1%}')
    return f'Bars sorted largest first with a cumulative % line. Running total of the top 3: {", ".join(cum)} → 3 causes ≈ 84% of defects.'


def make_start(ws, tricks):
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 3
    ws.column_dimensions['B'].width = 6
    ws.column_dimensions['C'].width = 44
    ws.column_dimensions['D'].width = 54
    ws.column_dimensions['E'].width = 22
    put(ws, 'B2', 'Excel to Excel · Practice workbook', bold=True, size=18, color=GREEN)
    put(ws, 'B3', '20 Excel tricks you don\'t know, for Industrial Engineers. Same data as the webinar slides.', size=11)
    put(ws, 'B5', 'How it works', bold=True, size=12)
    howto = [
        'Each sheet is one trick. The data sits in the same cells as on the slide, so the slide\'s formula works as-is.',
        'Yellow cells are yours: type there (or run the menu steps listed on the sheet).',
        'The Check cell turns green (✓) when you have it right, red (✗) when not yet.',
        'Stuck? Each sheet has a hidden ANSWER column group on the right: click its [+] button to open it.',
        'Example: on 07 SUMIFS type =SUMIFS(D2:D13,B2:B13,"L2",C2:C13,"Night") in G2 → Check shows ✓ Correct!',
    ]
    for i, line in enumerate(howto):
        put(ws, f'B{6 + i}', '•')
        put(ws, f'C{6 + i}', line)
    put(ws, 'B12', 'Legend:', bold=True)
    put(ws, 'C12', 'Type here', fill=YELLOW)
    ws['C12'].border = BOX
    put(ws, 'D12', '✓ Check cell: green when right, red when not yet', bold=True, color='008000')
    put(ws, 'B14', '#', bold=True, fill=HEAD)
    put(ws, 'C14', 'Trick (click to open)', bold=True, fill=HEAD)
    put(ws, 'D14', 'What you practise', bold=True, fill=HEAD)
    put(ws, 'E14', 'Works in', bold=True, fill=HEAD)
    for i, t in enumerate(tricks):
        r = 15 + i
        name = SHEETS[t['n']]
        put(ws, f'B{r}', t['n'])
        put(ws, f'C{r}', f'=HYPERLINK("#\'{name}\'!A1","{t["title"].replace(chr(34), "")}")', color='0563C1')
        put(ws, f'D{r}', t['tagline'])
        put(ws, f'E{r}', {'All': 'Excel 2013+', '2016+': 'Excel 2016+', '2019+': 'Excel 2019+', '365/2021': 'Microsoft 365 / 2021'}[t['version']])
    put(ws, f'B{16 + len(tricks)}', '=HYPERLINK("' + DECK_URL + '","Open the webinar deck: ' + DECK_URL + '")', color='0563C1')
    put(ws, f'B{17 + len(tricks)}', 'Example data is fictional, made for practice. Check your Excel version: File › Account › About Excel.',
        italic=True, color='595959')


def dump(ws):
    out = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            v = c.value
            if isinstance(v, (datetime.date, datetime.datetime)):
                v = (v - datetime.datetime(1899, 12, 30)).days if isinstance(v, datetime.datetime) else (v - datetime.date(1899, 12, 30)).days
            out[c.coordinate] = {'v': v, 't': c.data_type}
    return out


if __name__ == '__main__':
    build()
